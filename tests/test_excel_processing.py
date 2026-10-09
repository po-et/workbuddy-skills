"""Flat-table merge CLI regressions; all workbooks and caches are synthetic."""
from datetime import date, datetime
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = load_workbook = None

SCRIPT = Path(__file__).resolve().parents[1] / "skills/excel-processing/scripts/merge_xlsx.py"


@unittest.skipUnless(Workbook is not None, "requires existing openpyxl")
class ExcelMergeCLITest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="synthetic-excel-test-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.output = self.root / "combined.xlsx"

    def book(self, name, rows, sheet="明细", header_row=1):
        path = self.root / name
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for _ in range(header_row - 1):
            ws.append(["合成测试说明，不是明细"])
        for row in rows:
            ws.append(row)
        other = wb.create_sheet("不要合并")
        other.append(["其他表"])
        other.append(["=SUM(1,2)"])
        wb.save(path)
        wb.close()
        return path

    def run_cli(self, inputs, *args, select=True):
        hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs if path.exists()}
        command = [sys.executable, str(SCRIPT), *[str(path) for path in inputs], "--out", str(self.output)]
        if select:
            command += ["--sheet", "明细"]
        result = subprocess.run(command + list(args), capture_output=True, text=True,
                                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}, timeout=30)
        self.assertNotIn("Traceback", result.stderr)
        self.assertNotIn("Exception ignored", result.stderr)
        for path, digest in hashes.items():
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)
        return result

    def read_output(self):
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        return wb["合并明细"]

    def test_reordered_headers_dates_and_provenance(self):
        east = self.book("east source.xlsx", [["订单", "日期", "金额"], ["001", date(2026, 10, 1), 120],
                                            [None, None, None], ["002", date(2026, 10, 2), 150]])
        west = self.book("west.xlsx", [["金额", "订单", "日期"], [80, "003", date(2026, 10, 3)]])
        wb = load_workbook(east)
        wb["明细"]["C2"].number_format = "#,##0.00"
        wb.save(east)
        wb.close()
        result = self.run_cli([east, west])
        self.assertEqual(result.returncode, 0, result.stderr)
        summary = json.loads(result.stdout)
        self.assertEqual((summary["rows_read"], summary["rows_written"], summary["blank_rows_skipped"]), (4, 3, 1))
        ws = self.read_output()
        rows = list(ws.values)
        self.assertEqual(rows[0], ("订单", "日期", "金额", "来源文件", "来源工作表", "来源行"))
        self.assertEqual([row[0] for row in rows[1:]], ["001", "002", "003"])
        self.assertEqual([row[2] for row in rows[1:]], [120, 150, 80])
        self.assertIsInstance(rows[1][1], datetime)
        self.assertEqual(rows[1][3:], (str(east.resolve()), "明细", 2))
        self.assertEqual(rows[2][5], 4)
        self.assertEqual(ws["C2"].number_format, "#,##0.00")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:F4")

    def test_nonfirst_header_row_and_partial_empty_cells(self):
        path = self.book("prefix.xlsx", [["编号", "金额"], ["A1", None], ["A2", 0]], header_row=3)
        result = self.run_cli([path], "--header-row", "3")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(list(self.read_output().values)[1][:2], ("A1", None))
        self.assertEqual(list(self.read_output().values)[2][-1], 5)

    def test_duplicate_or_blank_headers_are_rejected(self):
        for headers in (["金额", " 金额 "], ["订单", None, "金额"], ["订单", "来源行"]):
            with self.subTest(headers=headers):
                path = self.book("bad-header.xlsx", [headers, ["A1", 80, 50]])
                result = self.run_cli([path])
                self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
                self.assertFalse(self.output.exists())

    def test_inconsistent_headers_fail_with_no_partial_output(self):
        east = self.book("east.xlsx", [["订单", "金额"], ["A1", 120]])
        west = self.book("west.xlsx", [["订单", "金额(元)"], ["A2", 80]])
        result = self.run_cli([east, west])
        self.assertEqual(result.returncode, 1)
        self.assertIn("缺少：金额；多出：金额(元)", result.stderr)
        self.assertFalse(self.output.exists())

    def test_explicit_sheet_is_required_and_not_guessed(self):
        path = self.book("other-sheet.xlsx", [["订单", "金额"], ["A1", 120]], sheet="其他明细")
        for select in (True, False):
            with self.subTest(select=select):
                result = self.run_cli([path], select=select)
                self.assertEqual(result.returncode, 1)
                self.assertFalse(self.output.exists())

    def test_existing_output_and_input_as_output_are_preserved(self):
        path = self.book("source.xlsx", [["订单", "金额"], ["A1", 120]])
        self.output.write_bytes(b"existing-output")
        result = self.run_cli([path])
        self.assertEqual(result.returncode, 2)
        self.assertEqual(self.output.read_bytes(), b"existing-output")
        self.output = path
        result = self.run_cli([path])
        self.assertEqual(result.returncode, 2)
        self.assertIn("原件", result.stderr)

    def test_formula_default_and_missing_cache_are_rejected(self):
        first = self.book("first.xlsx", [["订单", "金额"], ["A1", 120]])
        second = self.book("formula.xlsx", [["订单", "金额"], ["A2", "=SUM(2,3)"]])
        for args, message in (((), "发现公式"), (("--formulas", "cached"), "公式缓存缺失")):
            with self.subTest(args=args):
                result = self.run_cli([first, second], *args)
                self.assertEqual(result.returncode, 1, result.stderr)
                self.assertIn(message, result.stderr)
                self.assertIn("!B2", result.stderr)
                self.assertFalse(self.output.exists())

    def test_explicit_cached_mode_reads_synthetic_cache_as_value(self):
        path = self.book("cached-formula.xlsx", [["订单", "金额"], ["A1", "=SUM(2,3)"]])
        # A manufactured cache fixture proves reading behavior, not Excel recalculation.
        with zipfile.ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        tree = ET.fromstring(members["xl/worksheets/sheet1.xml"])
        tree.find('.//s:c[@r="B2"]/s:v', namespace).text = "5"
        members["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in members.items():
                archive.writestr(name, data)
        result = self.run_cli([path], "--formulas", "cached")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["cached_formula_cells"], 1)
        cell = self.read_output()["B2"]
        self.assertEqual(cell.value, 5)
        self.assertNotEqual(cell.data_type, "f")

    def test_literal_equals_strings_do_not_turn_into_formulas(self):
        path = self.book("literal.xlsx", [["编号", "说明"], ["A1", "=literal-text"]])
        wb = load_workbook(path)
        wb["明细"]["B2"].data_type = "s"
        wb.save(path)
        wb.close()
        result = self.run_cli([path])
        self.assertEqual(result.returncode, 0, result.stderr)
        cell = self.read_output()["B2"]
        self.assertEqual((cell.value, cell.data_type), ("=literal-text", "s"))

    def test_data_without_header_and_empty_tables_are_rejected(self):
        for rows in ([["编号"], ["A1", 80]], [["编号", "金额"]]):
            with self.subTest(rows=rows):
                path = self.book("empty-or-wide.xlsx", rows)
                result = self.run_cli([path])
                self.assertEqual(result.returncode, 1, result.stderr)
                self.assertFalse(self.output.exists())

    def test_duplicate_input_and_invalid_paths_are_rejected(self):
        path = self.book("source.xlsx", [["编号"], ["A1"]])
        for inputs, code in (([path, path], 1), ([self.root / "missing.xlsx"], 2)):
            with self.subTest(inputs=inputs):
                result = self.run_cli(inputs)
                self.assertEqual(result.returncode, code)
                self.assertFalse(self.output.exists())
        corrupted = self.root / "corrupt.xlsx"
        corrupted.write_bytes(b"invalid-zip")
        result = self.run_cli([corrupted])
        self.assertEqual(result.returncode, 1)

    def test_incorrect_stored_dimensions_do_not_drop_rows(self):
        path = self.book("wrong-dimensions.xlsx", [["编号", "金额"], ["A1", 120], ["A2", 80]])
        with zipfile.ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        tree = ET.fromstring(members["xl/worksheets/sheet1.xml"])
        tree.find("s:dimension", namespace).set("ref", "A1:A1")
        members["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in members.items():
                archive.writestr(name, data)
        result = self.run_cli([path])
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["rows_written"], 2)
        self.assertEqual([row[0] for row in list(self.read_output().values)[1:]], ["A1", "A2"])


if __name__ == "__main__":
    unittest.main()
