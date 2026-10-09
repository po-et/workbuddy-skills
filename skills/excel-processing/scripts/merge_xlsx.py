#!/usr/bin/env python3
"""Merge flat .xlsx tables by header name; inputs are never saved or modified.

Example: python3 merge_xlsx.py east.xlsx west.xlsx --sheet 明细 --out merged.xlsx
Requires existing Python 3.8+ and openpyxl. Sheet selection is explicit.
Formulas are rejected by default; --formulas cached uses previously saved values
and rejects missing caches. It never evaluates or copies formula references.
Outputs must be new .xlsx files. Values and number formats are copied, not styles,
charts, macros, merged cells or other workbook structure. Exit 0 success; 1 invalid
input/options; 2 filesystem/output problem; 3 openpyxl unavailable; 130 interrupted.
"""
import argparse
import json
import os
from pathlib import Path
import sys
import tempfile

PROVENANCE = ("来源文件", "来源工作表", "来源行")
MAX_ROWS = 1048576
MAX_COLUMNS = 16384


class MergeError(Exception):
    def __init__(self, message, code=1):
        super().__init__(message)
        self.code = code


class Parser(argparse.ArgumentParser):
    def error(self, message):
        raise MergeError("参数错误：" + message)


def dependencies():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise MergeError("缺少 openpyxl；请使用已具备该库的 Python 环境。本脚本不安装依赖。", 3)
    return Workbook, load_workbook, WriteOnlyCell, get_column_letter


def open_book(load_workbook, path, data_only=False):
    try:
        return load_workbook(path, read_only=True, data_only=data_only, keep_links=False)
    except Exception as exc:
        raise MergeError("无法读取 %s：%s" % (path, exc))


def selected_sheet(book, path, sheet_name):
    if sheet_name not in book.sheetnames:
        raise MergeError("%s 没有工作表「%s」；可用：%s" % (path, sheet_name, "、".join(book.sheetnames)))
    sheet = book[sheet_name]
    # A read-only workbook can have incorrect stored dimensions; inspect actual rows.
    sheet.reset_dimensions()
    return sheet


def header_names(sheet, path, header_row):
    row = next(sheet.iter_rows(min_row=header_row, max_row=header_row), ())
    values = [cell.value for cell in row]
    while values and values[-1] is None:
        values.pop()
    if not values:
        raise MergeError("%s/%s 第 %d 行没有表头" % (path, sheet.title, header_row))
    names = []
    for index, value in enumerate(values, 1):
        if not isinstance(value, str) or not value.strip():
            raise MergeError("%s/%s 第 %d 行第 %d 列表头必须为非空文本" % (path, sheet.title, header_row, index))
        if row[index - 1].data_type == "f":
            raise MergeError("%s/%s 表头不能是公式：%s" % (path, sheet.title, row[index - 1].coordinate))
        name = value.strip()
        if name in names:
            raise MergeError("%s/%s 重复表头「%s」（首尾空格已去除）" % (path, sheet.title, name))
        if name in PROVENANCE:
            raise MergeError("%s/%s 表头「%s」与输出追溯列冲突，请先重命名" % (path, sheet.title, name))
        names.append(name)
    if len(names) + len(PROVENANCE) > MAX_COLUMNS:
        raise MergeError("输出列数超过 Excel 上限 %d" % MAX_COLUMNS)
    return names


def merge(args):
    Workbook, load_workbook, WriteOnlyCell, get_column_letter = dependencies()
    inputs = [Path(value).expanduser().resolve() for value in args.inputs]
    output = Path(args.out).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        raise MergeError("--out 必须是新的 .xlsx 文件")
    if output in inputs:
        raise MergeError("输出不能是任一输入原件，请另存新文件", 2)
    if output.exists():
        raise MergeError("输出已存在，拒绝覆盖：%s；请换 --out" % output, 2)
    if len(set(inputs)) != len(inputs):
        raise MergeError("输入文件重复，请每个文件只指定一次")
    for path in inputs:
        if path.suffix.lower() != ".xlsx" or path.name.startswith("~$"):
            raise MergeError("只接受正常的 .xlsx 文件，不处理锁文件、.xls 或 .xlsm：%s" % path)
        if not path.is_file():
            raise MergeError("找不到输入文件：%s" % path, 2)

    columns, specs = None, []
    for path in inputs:
        book = open_book(load_workbook, path)
        try:
            sheet = selected_sheet(book, path, args.sheet)
            names = header_names(sheet, path, args.header_row)
            if columns is None:
                columns = names
            elif set(names) != set(columns):
                missing = [name for name in columns if name not in names]
                extra = [name for name in names if name not in columns]
                raise MergeError("%s/%s 表头不一致；缺少：%s；多出：%s" %
                                 (path, sheet.title, "、".join(missing) or "无", "、".join(extra) or "无"))
            specs.append((path, names))
        finally:
            book.close()

    out_book = Workbook(write_only=True)
    out_sheet = out_book.create_sheet("合并明细")
    out_sheet.freeze_panes = "A2"
    # Force header text, including names starting with '=', to remain text.
    header_cells = []
    for value in columns + list(PROVENANCE):
        cell = WriteOnlyCell(out_sheet, value=value)
        cell.data_type = "s"
        header_cells.append(cell)
    out_sheet.append(header_cells)
    summary = {"output": str(output), "sheet": "合并明细", "columns": columns + list(PROVENANCE),
               "formulas": args.formulas, "rows_read": 0, "rows_written": 0,
               "blank_rows_skipped": 0, "cached_formula_cells": 0, "sources": []}
    temporary = None
    try:
        for path, names in specs:
            book, cache_book = open_book(load_workbook, path), None
            source = {"file": str(path), "sheet": args.sheet, "rows_read": 0,
                      "rows_written": 0, "blank_rows_skipped": 0}
            try:
                sheet = selected_sheet(book, path, args.sheet)
                rows = sheet.iter_rows(min_row=args.header_row + 1)
                cache_rows = None
                if args.formulas == "cached":
                    cache_book = open_book(load_workbook, path, data_only=True)
                    cache_rows = selected_sheet(cache_book, path, args.sheet).iter_rows(min_row=args.header_row + 1)
                positions = [names.index(name) for name in columns]
                for row_number, row in enumerate(rows, args.header_row + 1):
                    cached = next(cache_rows, ()) if cache_rows is not None else ()
                    source["rows_read"] += 1
                    if not any(cell.value is not None for cell in row):
                        source["blank_rows_skipped"] += 1
                        continue
                    if any(cell.value is not None for cell in row[len(names):]):
                        raise MergeError("%s/%s 第 %d 行有数据位于无表头列，拒绝丢弃" % (path, sheet.title, row_number))
                    result = []
                    for position in positions:
                        original = row[position] if position < len(row) else None
                        selected = original
                        if original is not None and original.data_type == "f":
                            location = "%s/%s!%s" % (path, sheet.title, original.coordinate)
                            if args.formulas == "reject":
                                raise MergeError("发现公式 %s；默认拒绝迁移引用。先在 Excel 中计算保存并确认缓存，再明确选 --formulas cached" % location)
                            selected = cached[position] if position < len(cached) else None
                            if selected is None or selected.value is None or selected.data_type == "e":
                                raise MergeError("公式缓存缺失或为错误值：%s；openpyxl 不计算公式，请先用 Excel 计算保存" % location)
                            summary["cached_formula_cells"] += 1
                        cell = WriteOnlyCell(out_sheet, value=selected.value if selected is not None else None)
                        if selected is not None:
                            cell.number_format = original.number_format or "General"
                            # Literal strings beginning with '=' must not become new formulas.
                            if selected.data_type in ("s", "inlineStr"):
                                cell.data_type = "s"
                        result.append(cell)
                    for value in (str(path), args.sheet):
                        cell = WriteOnlyCell(out_sheet, value=value)
                        cell.data_type = "s"
                        result.append(cell)
                    result.append(row_number)
                    if summary["rows_written"] + 1 >= MAX_ROWS:
                        raise MergeError("输出超过 Excel 单表行数上限 %d（含表头）" % MAX_ROWS)
                    out_sheet.append(result)
                    source["rows_written"] += 1
                    summary["rows_written"] += 1
                summary["rows_read"] += source["rows_read"]
                summary["blank_rows_skipped"] += source["blank_rows_skipped"]
                summary["sources"].append(source)
            finally:
                book.close()
                if cache_book is not None:
                    cache_book.close()
        if not summary["rows_written"]:
            raise MergeError("所选工作表没有可合并的数据行，不生成空结果")
        out_sheet.auto_filter.ref = "A1:%s%d" % (
            get_column_letter(len(columns) + len(PROVENANCE)), summary["rows_written"] + 1)
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix=".merge-xlsx-", suffix=".xlsx", dir=str(output.parent))
        os.close(fd)
        temporary = Path(name)
        out_book.save(temporary)
        # Atomic no-clobber publication; an existing path cannot be overwritten even in a race.
        os.link(str(temporary), str(output))
        return summary
    except OSError as exc:
        raise MergeError("无法写入新输出 %s：%s" % (output, exc), 2)
    finally:
        if not out_sheet.closed:
            out_sheet.close()
        out_book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main(argv=None):
    parser = Parser(description="按表头名合并平面 .xlsx 明细表；原件只读，结果另存且拒绝覆盖")
    parser.add_argument("inputs", nargs="+", help="输入 .xlsx 文件，按给定顺序合并")
    parser.add_argument("--sheet", required=True, help="每个输入中要选的确切工作表名，不自动遍历所有表")
    parser.add_argument("--header-row", type=int, default=1, help="表头行号（默认 1，所有输入一致）")
    parser.add_argument("--formulas", choices=("reject", "cached"), default="reject", help="公式处理：默认拒绝；cached 仅复制非空已有缓存值")
    parser.add_argument("--out", required=True, help="新的 .xlsx 输出路径，不提供覆盖选项")
    try:
        args = parser.parse_args(argv)
        if args.header_row < 1 or args.header_row >= MAX_ROWS:
            raise MergeError("--header-row 必须在 1 到 %d 之间" % (MAX_ROWS - 1))
        result = merge(args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except MergeError as exc:
        print(str(exc), file=sys.stderr)
        return exc.code


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("已中断；输入原件未保存或修改，请检查是否已有输出后换路径重跑。", file=sys.stderr)
        sys.exit(130)
